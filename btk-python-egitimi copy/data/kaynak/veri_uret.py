# -*- coding: utf-8 -*-
"""Eğitim için sentetik telekom verisi üretir (gerçek kişi/kurum/operatör verisi DEĞİLDİR).

Çalıştırma (paketli ortam, openpyxl gerekir): python data/kaynak/veri_uret.py
Sabit tohum (seed) kullanıldığı için her çalıştırmada AYNI veri üretilir.
"""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

DATA = Path(__file__).resolve().parents[1]
rnd = random.Random(2026)

ILLER = ["Ankara", "İstanbul", "İzmir", "Bursa", "Antalya", "Konya", "Erzurum", "Trabzon"]
IL_AGIRLIK = [14, 30, 12, 9, 8, 7, 5, 5]
OPERATORLER = ["Operatör A", "Operatör B", "Operatör C", "Operatör D"]
OP_AGIRLIK = [38, 30, 22, 10]
KATEGORILER = ["Fatura", "İnternet Hızı", "Kesinti", "Taahhüt/İptal", "Numara Taşıma", "Kapsama"]
KAT_AGIRLIK = [26, 22, 18, 16, 8, 10]
KANALLAR = ["Web", "Çağrı Merkezi", "e-Devlet", "Dilekçe"]
KANAL_AGIRLIK = [35, 25, 32, 8]
HIZMET = ["Mobil", "Sabit İnternet", "Sabit Telefon"]

# Kategoriye göre ortalama çözüm süresi (gün) — analizde "anlamlı fark" görülsün diye
KAT_SURE = {"Fatura": 6, "İnternet Hızı": 11, "Kesinti": 4, "Taahhüt/İptal": 14,
            "Numara Taşıma": 3, "Kapsama": 18}
OP_CARPAN = {"Operatör A": 1.0, "Operatör B": 0.85, "Operatör C": 1.25, "Operatör D": 1.1}

baslangic = date(2026, 1, 1)
sikayetler = []
for no in range(1, 601):
    kat = rnd.choices(KATEGORILER, KAT_AGIRLIK)[0]
    op = rnd.choices(OPERATORLER, OP_AGIRLIK)[0]
    # Yaz aylarında kesinti şikâyeti biraz artsın (mevsimsellik)
    gun_ofset = rnd.randint(0, 272)  # 1 Ocak - 30 Eylül 2026
    tarih = baslangic + timedelta(days=gun_ofset)
    if kat == "Kesinti" and rnd.random() < 0.35:
        tarih = date(2026, rnd.choice([6, 7, 8]), rnd.randint(1, 28))
    hizmet = ("Mobil" if kat in ("Kapsama", "Numara Taşıma") else
              rnd.choices(HIZMET, [45, 45, 10])[0])
    sure = max(1, round(rnd.gauss(KAT_SURE[kat] * OP_CARPAN[op], 3)))
    durum = rnd.choices(["Çözüldü", "Reddedildi", "İşlemde"], [70, 12, 18])[0]
    memnuniyet = ""
    if durum == "Çözüldü":
        # Uzun süren şikâyette memnuniyet düşer
        taban = 5 - sure / 6 + rnd.uniform(-1, 1)
        memnuniyet = min(5, max(1, round(taban)))
    elif durum == "Reddedildi":
        memnuniyet = rnd.choice([1, 1, 2, 2, 3])
    tutar = round(rnd.uniform(50, 2500), 2) if kat == "Fatura" else ""
    sikayetler.append({
        "sikayet_no": 10000 + no,
        "tarih": tarih.isoformat(),
        "il": rnd.choices(ILLER, IL_AGIRLIK)[0],
        "operator": op,
        "hizmet": hizmet,
        "kategori": kat,
        "kanal": rnd.choices(KANALLAR, KANAL_AGIRLIK)[0],
        "durum": durum,
        "cozum_suresi_gun": sure if durum != "İşlemde" else "",
        "memnuniyet": memnuniyet,
        "itiraz_tutari_tl": tutar,
    })
sikayetler.sort(key=lambda r: (r["tarih"], r["sikayet_no"]))

TEKNOLOJI = {"Fiber": (180, 60), "VDSL": (45, 12), "ADSL": (12, 4), "4.5G": (75, 30)}
olcumler = []
for no in range(1, 401):
    tek = rnd.choices(list(TEKNOLOJI), [35, 30, 10, 25])[0]
    ort, sapma = TEKNOLOJI[tek]
    paket = {"Fiber": rnd.choice([100, 200, 500, 1000]), "VDSL": rnd.choice([35, 50, 100]),
             "ADSL": rnd.choice([8, 16, 24]), "4.5G": rnd.choice([50, 100, 150])}[tek]
    indirme = max(0.5, min(paket * 1.02, rnd.gauss(min(ort, paket * 0.82), sapma)))
    olcumler.append({
        "olcum_no": no,
        "tarih": (baslangic + timedelta(days=rnd.randint(0, 272))).isoformat(),
        "il": rnd.choices(ILLER, IL_AGIRLIK)[0],
        "operator": rnd.choices(OPERATORLER, OP_AGIRLIK)[0],
        "teknoloji": tek,
        "paket_hizi_mbps": paket,
        "indirme_mbps": round(indirme, 1),
        "yukleme_mbps": round(indirme * (0.5 if tek == "Fiber" else 0.15) * rnd.uniform(0.7, 1.1), 1),
        "ping_ms": round(rnd.uniform(4, 15) if tek == "Fiber" else rnd.uniform(12, 60), 1),
    })


def yaz(ad, satirlar):
    with (DATA / ad).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(satirlar[0]))
        w.writeheader()
        w.writerows(satirlar)


yaz("sikayetler.csv", sikayetler)
yaz("hiz_olcumleri.csv", olcumler)

# Kasıtlı kirli küçük örnek: temizleme dersleri için
kirli = [dict(r) for r in sikayetler[:25]]
kirli[2]["il"] = " ankara "
kirli[4]["operator"] = "operatör a"
kirli[6]["cozum_suresi_gun"] = "on iki"
kirli[8]["tarih"] = "31.02.2026"
kirli[10]["memnuniyet"] = "7"
kirli[12]["cozum_suresi_gun"] = "-3"
kirli.append(dict(kirli[1]))  # tekrar eden satır
kirli[15]["kategori"] = ""
yaz("sikayetler_kirli.csv", kirli)

try:
    from openpyxl import Workbook
except ImportError:
    print("openpyxl yok: .xlsx üretilmedi (CSV dosyaları hazır).")
else:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sikayetler"
    ws.append(list(sikayetler[0]))
    for r in sikayetler:
        ws.append([r[k] if r[k] != "" else None for k in r])
    ws2 = wb.create_sheet("Hiz_Olcumleri")
    ws2.append(list(olcumler[0]))
    for r in olcumler:
        ws2.append(list(r.values()))
    ws3 = wb.create_sheet("Hedefler")
    ws3.append(["kategori", "hedef_gun"])
    for k, v in {"Fatura": 7, "İnternet Hızı": 10, "Kesinti": 5, "Taahhüt/İptal": 15,
                 "Numara Taşıma": 3, "Kapsama": 20}.items():
        ws3.append([k, v])
    wb.save(DATA / "btk_ornek_veri.xlsx")
    print("btk_ornek_veri.xlsx yazıldı (3 sayfa).")
print(f"sikayetler.csv: {len(sikayetler)} satır, hiz_olcumleri.csv: {len(olcumler)} satır")
